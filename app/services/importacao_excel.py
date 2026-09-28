# -*- coding: utf-8 -*-
"""Importação dos exports do sistema de vendas para dentro do sistema.

Lê `Export_F10.xlsx` (cadastro) e `Export_F10 - Chamadas.xlsx` (presenças) e
SUBSTITUI todos os dados atuais do banco. Antes de apagar, gera um backup do
banco atual em `backups/`.

Usada pela tela de Importação do sistema (rota `/importacao`) e pelo script
de linha de comando `importar_excel.py`. A função `importar()` deve ser chamada
dentro de um app context (rotas Flask já oferecem; o CLI cria o app).
"""
import re
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl

from config import BACKUPS_DIR, DATA_DIR
from extensions import db
from models import (
    AUSENTE,
    FREQUENTE,
    Aluno,
    Chamada,
    Contato,
    Curso,
    Materia,
    Matricula,
    Presenca,
    Turma,
)
from services.materias import MATERIAS_PADRAO_INFORMATICA, MATERIA_NEUTRA

NOME_CURSO = "Informatica"
CONTRATO_ATIVO = "Ativo"

DIAS = ["segunda", "terca", "quarta", "quinta", "sexta", "sabado", "domingo"]


def limpar_telefone(valor):
    if not isinstance(valor, str):
        return None
    texto = valor.strip()
    if not texto or texto.lower().startswith("n"):
        return None
    texto = re.split(r"[;|]", texto)[0].strip()
    return texto or None


def limpar_fio_nome(valor):
    if not isinstance(valor, str):
        return None
    texto = valor.strip()
    return texto or None


def ler_cadastro(caminho):
    """Retorna: nome_da_turma -> lista de dicts (linha por aluno/turma)."""
    wb = openpyxl.load_workbook(caminho, data_only=True)
    saida = {}
    for nome in wb.sheetnames:
        ws = wb[nome]
        linhas = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r[0]:
                continue
            linhas.append(
                {
                    "aluno": r[0].strip(),
                    "titular": limpar_fio_nome(r[1]),
                    "codigo_turma": (
                        (r[4] or "").strip()
                        if isinstance(r[4], str)
                        else r[4]
                    ),
                    "entrada": r[9],
                    "tel_aluno": r[10],
                    "cel_aluno": r[11],
                    "tel_titular": r[13],
                    "cel_titular": r[14],
                    "com_titular": r[15],
                    "com_aluno": r[16],
                    "status_contrato": (
                        (r[3] or "").strip()
                        if isinstance(r[3], str)
                        else r[3]
                    ),
                }
            )
        saida[nome] = linhas
    wb.close()
    return saida


def ler_chamadas(caminho):
    """Retorna: nome_da_turma -> {'datas': [date...], 'notas': {aluno: {date: 0|1}}}."""
    wb = openpyxl.load_workbook(caminho, data_only=True)
    saida = {}
    for nome in wb.sheetnames:
        ws = wb[nome]
        linhas = list(ws.iter_rows(values_only=True))
        hdr = linhas[0]
        datas = []
        for h in hdr[1:]:
            if not isinstance(h, datetime):
                continue
            d = h.date()
            if d not in datas:
                datas.append(d)
        notas = {}
        for r in linhas[1:]:
            if not r or not r[0]:
                continue
            nome_aluno = r[0].strip()
            if nome_aluno in notas:
                continue
            cel = {
                datas[j]: (r[1 + j] if len(r) > 1 + j else None)
                for j in range(len(datas))
            }
            notas[nome_aluno] = cel
        saida[nome] = {"datas": datas, "notas": notas}
    wb.close()
    return saida


def telefones_do_aluno(linha):
    if linha["titular"]:
        return (
            limpar_telefone(linha["tel_titular"]),
            limpar_telefone(linha["cel_titular"]),
            limpar_telefone(linha["com_titular"]),
        )
    return (
        limpar_telefone(linha["tel_aluno"]),
        limpar_telefone(linha["cel_aluno"]),
        limpar_telefone(linha["com_aluno"]),
    )


def backup_atual():
    """Copia o banco atual para backups/ (se existir) e devolve o nome."""
    origem = DATA_DIR / "sistema.db"
    if not origem.exists():
        return None
    destino = BACKUPS_DIR / (
        "sistema_pre_importacao_" + datetime.now().strftime("%Y-%m-%d_%H%M%S")
        + ".db"
    )
    shutil.copy2(origem, destino)
    return destino.name


def importar(caminho_cadastro, caminho_chamadas):
    """Lê os exports e substitui os dados do banco (requer app context).

    Retorna dict com o resumo da importação. Em erro, faz rollback e relança.
    """
    cad = ler_cadastro(caminho_cadastro)
    cha = ler_chamadas(caminho_chamadas)

    if not cad:
        raise ValueError("O arquivo de cadastro não trouxe nenhuma turma.")

    nome_backup = backup_atual()

    try:
        db.session.query(Presenca).delete()
        db.session.query(Matricula).delete()
        db.session.query(Chamada).delete()
        db.session.query(Materia).delete()
        db.session.query(Contato).delete()
        db.session.query(Aluno).delete()
        db.session.query(Turma).delete()
        db.session.query(Curso).delete()

        curso = Curso(nome=NOME_CURSO)
        db.session.add(curso)
        db.session.flush()

        # O curso importado já nasce com a "Neutra" (+ padrão Informática),
        # para a matriz e o lançamento funcionarem sem depender de restart.
        neutra = Materia(curso_id=curso.id, nome=MATERIA_NEUTRA)
        db.session.add(neutra)
        for nome_materia in MATERIAS_PADRAO_INFORMATICA:
            db.session.add(
                Materia(curso_id=curso.id, nome=nome_materia)
            )
        db.session.flush()

        turmas_por_sheet = {}
        turmas_por_codigo = {}
        for nome_sheet in cad:
            datas = cha.get(nome_sheet, {}).get("datas", [])
            dia_semana = DIAS[datas[0].weekday()] if datas else "segunda"
            codigo = None
            for linha in cad[nome_sheet]:
                if linha["codigo_turma"]:
                    codigo = linha["codigo_turma"]
                    break
            nome_turma = codigo or nome_sheet
            turma = turmas_por_codigo.get(codigo) if codigo else None
            if turma is None:
                turma = Turma(
                    curso_id=curso.id,
                    nome=nome_turma,
                    dia_semana=dia_semana,
                    ativa=True,
                )
                db.session.add(turma)
                db.session.flush()
                if codigo:
                    turmas_por_codigo[codigo] = turma
            turmas_por_sheet[nome_sheet] = turma

        alunos_por_nome = {}
        for nome_sheet, linhas in cad.items():
            for linha in linhas:
                nome = linha["aluno"]
                aluno = alunos_por_nome.get(nome)
                tel, cel, com = telefones_do_aluno(linha)
                entrada = linha["entrada"]
                if aluno is None:
                    aluno = Aluno(
                        nome=nome,
                        responsavel=linha["titular"],
                        telefone=tel,
                        celular=cel,
                        comercial=com,
                        data_matricula=(
                            entrada
                            if isinstance(entrada, datetime)
                            else datetime.today().date()
                        ),
                        status="inativo",
                    )
                    db.session.add(aluno)
                    db.session.flush()
                    alunos_por_nome[nome] = aluno
                else:
                    if aluno.responsavel is None and linha["titular"]:
                        aluno.responsavel = linha["titular"]
                    if aluno.telefone is None and tel:
                        aluno.telefone = tel
                    if aluno.celular is None and cel:
                        aluno.celular = cel
                    if aluno.comercial is None and com:
                        aluno.comercial = com
                    data_inicio = (
                        entrada if isinstance(entrada, datetime) else None
                    )
                    if data_inicio and (
                        aluno.data_matricula is None
                        or data_inicio < aluno.data_matricula
                    ):
                        aluno.data_matricula = data_inicio
                if linha["status_contrato"] == CONTRATO_ATIVO:
                    aluno.status = "ativo"

        matriculas_ativas = {}
        matriculas_por_par = {}
        for nome_sheet, linhas in cad.items():
            turma = turmas_por_sheet[nome_sheet]
            for linha in linhas:
                aluno = alunos_por_nome[linha["aluno"]]
                par = (turma.id, aluno.id)
                ativa = linha["status_contrato"] == CONTRATO_ATIVO
                matricula = matriculas_por_par.get(par)
                if matricula is None:
                    matricula = Matricula(
                        aluno_id=aluno.id,
                        turma_id=turma.id,
                        data_inicio=(
                            linha["entrada"]
                            if isinstance(linha["entrada"], datetime)
                            else None
                        ),
                        ativa=ativa,
                    )
                    db.session.add(matricula)
                    db.session.flush()
                    matriculas_por_par[par] = matricula
                else:
                    data_inicio = (
                        linha["entrada"]
                        if isinstance(linha["entrada"], datetime)
                        else None
                    )
                    if data_inicio and matricula.data_inicio is None:
                        matricula.data_inicio = data_inicio
                    if ativa and not matricula.ativa:
                        matricula.ativa = True
                if matriculas_por_par[par].ativa:
                    matriculas_ativas[par] = matriculas_por_par[par]

        chamadas_por_sheet = {}
        chamadas_feitas = set()
        for nome_sheet, dados in cha.items():
            turma = turmas_por_sheet[nome_sheet]
            for data in dados["datas"]:
                chave = (turma.id, data)
                if chave in chamadas_feitas:
                    continue
                chamadas_feitas.add(chave)
                chamada = Chamada(
                    turma_id=turma.id,
                    materia_id=neutra.id,
                    data=data,
                )
                db.session.add(chamada)
                db.session.flush()
                chamadas_por_sheet.setdefault(nome_sheet, []).append(chamada)

        for nome_sheet, dados in cha.items():
            turma = turmas_por_sheet[nome_sheet]
            for chamada in chamadas_por_sheet.get(nome_sheet, []):
                for nome_aluno, cel in dados["notas"].items():
                    aluno = alunos_por_nome.get(nome_aluno)
                    if aluno is None:
                        continue
                    valor = cel.get(chamada.data)
                    if (turma.id, aluno.id) in matriculas_ativas:
                        estado = FREQUENTE if valor == 1 else AUSENTE
                        db.session.add(
                            Presenca(
                                chamada_id=chamada.id,
                                aluno_id=aluno.id,
                                estado=estado,
                            )
                        )
                    elif valor == 1:
                        db.session.add(
                            Presenca(
                                chamada_id=chamada.id,
                                aluno_id=aluno.id,
                                estado=AUSENTE,
                            )
                        )

        db.session.commit()
    except Exception:
        db.session.rollback()
        raise

    total_alunos = len(alunos_por_nome)
    total_ativos = sum(
        1 for a in alunos_por_nome.values() if a.status == "ativo"
    )
    return {
        "turmas": len(turmas_por_sheet),
        "alunos": total_alunos,
        "alunos_ativos": total_ativos,
        "matriculas": Matricula.query.count(),
        "chamadas": Chamada.query.count(),
        "presencas": Presenca.query.count(),
        "backup": nome_backup,
    }